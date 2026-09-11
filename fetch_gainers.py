"""Fetch Upbit KRW market top-10 gainers and append them to coin_gainers.xlsx."""
import sys
from datetime import datetime, timedelta, timezone

import requests
from openpyxl import Workbook, load_workbook

MARKET_URL = "https://api.upbit.com/v1/market/all?is_details=false"
TICKER_URL = "https://api.upbit.com/v1/ticker"
XLSX_PATH = "coin_gainers.xlsx"
HEADER = ["실행시각(KST)", "순위", "마켓코드", "코인명", "등락률(%)", "현재가"]
KST = timezone(timedelta(hours=9))


def fetch_krw_markets():
    resp = requests.get(MARKET_URL, timeout=10)
    resp.raise_for_status()
    markets = resp.json()
    return {
        m["market"]: m["korean_name"]
        for m in markets
        if m["market"].startswith("KRW-")
    }


def fetch_tickers(market_codes):
    tickers = []
    # Upbit limits how many markets can be queried per request; chunk to be safe.
    chunk_size = 100
    for i in range(0, len(market_codes), chunk_size):
        chunk = market_codes[i : i + chunk_size]
        resp = requests.get(
            TICKER_URL, params={"markets": ",".join(chunk)}, timeout=10
        )
        resp.raise_for_status()
        tickers.extend(resp.json())
    return tickers


def main():
    try:
        market_to_name = fetch_krw_markets()
        tickers = fetch_tickers(list(market_to_name.keys()))
    except requests.RequestException as e:
        print(f"Upbit API 호출 실패: {e}", file=sys.stderr)
        sys.exit(1)

    if not tickers:
        print("Upbit API 응답에 데이터가 없습니다.", file=sys.stderr)
        sys.exit(1)

    top10 = sorted(tickers, key=lambda t: t["signed_change_rate"], reverse=True)[:10]

    now_kst = datetime.now(KST).strftime("%Y-%m-%d %H:%M:%S")

    try:
        wb = load_workbook(XLSX_PATH)
        ws = wb.active
    except FileNotFoundError:
        wb = Workbook()
        ws = wb.active
        ws.append(HEADER)

    for rank, t in enumerate(top10, start=1):
        market = t["market"]
        change_rate_pct = round(t["signed_change_rate"] * 100, 2)
        ws.append(
            [
                now_kst,
                rank,
                market,
                market_to_name.get(market, ""),
                change_rate_pct,
                t["trade_price"],
            ]
        )

    wb.save(XLSX_PATH)
    print(f"{now_kst} 기준 TOP10 저장 완료")


if __name__ == "__main__":
    main()
